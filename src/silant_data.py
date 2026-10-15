"""Загрузка, приведение к long-формату и проверки качества рыночных данных СИЛАНТ.

Парсер не завязан на фиксированные номера строк/столбцов: блоки ищутся по заголовку
(«Автопогрузчиков» / «Электропогрузчиков»), кварталы — по строке-шапке. Поэтому при
добавлении в файл нового квартала (новый столбец справа) весь пайплайн пересчитывается
без правок кода.
"""
from __future__ import annotations

import re
from pathlib import Path

import numpy as np
import openpyxl
import pandas as pd

BLOCKS = {"Автопогрузчик": "ДВС", "Электропогрузчик": "Электро"}
Q_RE = re.compile(r"(\d)\s*кв\.?\s*(\d{4})")
NUM_RE = re.compile(r"\d+(?:[.,]\d+)?")

# Гармонизированные классы грузоподъёмности для сравнения ДВС и электро
# (у типов разные «хвосты»: ДВС 6–8 / 8,5–10 / 11–18 / 20–25 т, электро 5,5–8 / 8,5–16 т).


def capacity_class(cap_min: float) -> str:
    """Map a minimum load capacity (t) to the class label used in the case."""
    if cap_min < 2:
        return "1–1,8 т"
    if cap_min < 3:
        return "2–2,5 т"
    if cap_min < 3.8:
        return "3–3,5 т"
    if cap_min < 5.5:
        return "3,8–5 т"
    if cap_min < 8.5:
        return "5,5–8 т"
    return "8,5 т и выше"


CLASS_ORDER = ["1–1,8 т", "2–2,5 т", "3–3,5 т", "3,8–5 т", "5,5–8 т", "8,5 т и выше"]


def _parse_capacity(label: str) -> tuple[float, float]:
    nums = [float(x.replace(",", ".")) for x in NUM_RE.findall(str(label))]
    return (nums[0], nums[-1]) if nums else (np.nan, np.nan)


def _seg_label(cmin: float, cmax: float) -> str:
    f = lambda v: (f"{v:g}").replace(".", ",")
    return f"{f(cmin)}–{f(cmax)} т"


def load_market(path: str | Path) -> tuple[pd.DataFrame, pd.DataFrame, list[dict]]:
    """Возвращает (long, totals_from_file, formula_info).

    long: type, segment, cap_min, cap_max, capacity_class, power, quarter, date, units
    totals_from_file: type, quarter, total_file — значения строки «Итого» (кэш Excel)
    formula_info: исходные формулы «Итого» для проверки диапазонов
    """
    path = Path(path)
    ws_f = openpyxl.load_workbook(path, data_only=False).worksheets[0]
    ws_v = openpyxl.load_workbook(path, data_only=True).worksheets[0]
    rows_f = [list(r) for r in ws_f.iter_rows(values_only=True)]
    rows_v = [list(r) for r in ws_v.iter_rows(values_only=True)]

    records, totals, formulas = [], [], []
    i = 0
    while i < len(rows_v):
        text = " ".join(str(c) for c in rows_v[i] if c is not None)
        block = next((t for k, t in BLOCKS.items() if k in text and "Объем" in text), None)
        if block is None:
            i += 1
            continue
        header = rows_v[i + 1]
        qcols = {j: Q_RE.search(str(h)) for j, h in enumerate(header) if h and Q_RE.search(str(h))}
        qcols = {j: pd.Period(f"{m.group(2)}Q{m.group(1)}", freq="Q") for j, m in qcols.items()}
        j = i + 2
        while j < len(rows_v) and rows_v[j][0] is not None:
            label = str(rows_v[j][0]).strip()
            if label.lower().startswith("итого"):
                for c, q in qcols.items():
                    totals.append({"type": block, "quarter": q, "total_file": rows_v[j][c]})
                formulas.append({"type": block, "excel_row": j + 1,
                                 "formula_first_col": rows_f[j][min(qcols)]})
                break
            cmin, cmax = _parse_capacity(label)
            for c, q in qcols.items():
                records.append({
                    "type": block, "segment_raw": label, "cap_min": cmin, "cap_max": cmax,
                    "power": rows_v[j][1], "quarter": q, "units": rows_v[j][c],
                    "excel_row": j + 1,
                })
            j += 1
        i = j + 1

    long = pd.DataFrame(records)
    long["segment"] = [_seg_label(a, b) for a, b in zip(long.cap_min, long.cap_max)]
    long["capacity_class"] = long.cap_min.map(capacity_class)
    long["date"] = long.quarter.dt.to_timestamp(how="end").dt.normalize()
    long["year"] = long.quarter.dt.year
    long["q"] = long.quarter.dt.quarter
    long["units"] = pd.to_numeric(long["units"], errors="coerce")
    cols = ["type", "segment", "cap_min", "cap_max", "capacity_class", "power",
            "quarter", "date", "year", "q", "units", "segment_raw", "excel_row"]
    long = long[cols].sort_values(["type", "cap_min", "quarter"]).reset_index(drop=True)
    return long, pd.DataFrame(totals), formulas


# ---------------------------------------------------------------- проверки
def check_totals(long: pd.DataFrame, totals: pd.DataFrame) -> pd.DataFrame:
    """Compare the sum of segments with the "Total" row of the customer file."""
    s = long.groupby(["type", "quarter"], as_index=False)["units"].sum()
    out = s.merge(totals, on=["type", "quarter"], how="outer")
    out["diff"] = out["units"] - pd.to_numeric(out["total_file"], errors="coerce")
    out["ok"] = out["diff"].fillna(np.inf).abs() < 1e-9
    return out.rename(columns={"units": "sum_segments"})


def check_formula_ranges(long: pd.DataFrame, formulas: list[dict]) -> pd.DataFrame:
    """Check that the SUM formulas of the "Total" rows cover exactly the segment rows."""
    rows = []
    for f in formulas:
        m = re.search(r"SUM\([A-Z]+(\d+):[A-Z]+(\d+)\)", str(f["formula_first_col"]))
        seg_rows = long.loc[long.type == f["type"], "excel_row"]
        exp = (seg_rows.min(), seg_rows.max())
        got = (int(m.group(1)), int(m.group(2))) if m else (None, None)
        rows.append({"type": f["type"], "formula": f["formula_first_col"],
                     "segment_rows": f"{exp[0]}–{exp[1]}", "formula_rows": f"{got[0]}–{got[1]}",
                     "exact_match": got == exp,
                     # лишние строки только сверху (название блока и шапка) — SUM их не учитывает
                     "header_only_extra": bool(m) and got[1] == exp[1] and exp[0] - 2 <= got[0] < exp[0]})
    return pd.DataFrame(rows)


def basic_quality(long: pd.DataFrame) -> pd.DataFrame:
    """Per-segment counts of missing, negative and zero values and a multiple-of-10 check."""
    g = long.groupby(["type", "segment"], sort=False)["units"]
    return pd.DataFrame({
        "n_quarters": g.size(),
        "n_missing": g.apply(lambda s: s.isna().sum()),
        "n_negative": g.apply(lambda s: (s < 0).sum()),
        "n_zero": g.apply(lambda s: (s == 0).sum()),
        "all_multiple_of_10": g.apply(lambda s: bool((s % 10 == 0).all())),
        "mean": g.mean().round(0),
        "min": g.min(), "max": g.max(),
        "cv": (g.std() / g.mean()).round(2),
    }).reset_index()


def detect_anomalies(long: pd.DataFrame, min_median: float = 100, t_thr: float = 3.0) -> pd.DataFrame:
    """Кандидаты в аномалии: внешне стьюдентизированные остатки модели
    log(y) = уровень + линейный тренд + сезонные дамми (по каждому сегменту отдельно).

    Стьюдентизированный остаток точки считается по модели, оценённой БЕЗ этой точки,
    поэтому выброс не «прячет» сам себя. На 13 точках (5 параметров, 8 степеней свободы)
    это фильтр «что проверить глазами», а не формальный тест.
    Сегменты с медианой < min_median шт. не проверяются: колебания 0↔40 шт. — норма
    для штучного спроса, для них отдельно считаем долю нулей.
    """
    import statsmodels.api as sm
    out = []
    for (t, seg), d in long.groupby(["type", "segment"], sort=False):
        d = d.sort_values("quarter")
        if d.units.median() < min_median:
            continue
        y = np.log(d.units.clip(lower=1).to_numpy(float))
        n = len(y)
        X = np.column_stack([np.ones(n), np.arange(n)] +
                            [(d.q.to_numpy() == k).astype(float) for k in (2, 3, 4)])
        fit = sm.OLS(y, X).fit()
        st = fit.get_influence().resid_studentized_external
        for q, u, f, z in zip(d.quarter, d.units, fit.fittedvalues, st):
            out.append({"type": t, "segment": seg, "quarter": q, "units": u,
                        "expected": int(round(float(np.exp(f)), -1)),
                        "stud_resid": round(float(z), 2)})
    a = pd.DataFrame(out)
    a["dev_pct"] = (a.units / a.expected - 1).round(3)
    a["flag"] = a.stud_resid.abs() >= t_thr
    return a


def yoy(long: pd.DataFrame, by=("type", "segment")) -> pd.DataFrame:
    """Year-over-year change of sales by the given grouping."""
    w = long.groupby(list(by) + ["quarter"])["units"].sum().unstack("quarter")
    return w.pct_change(periods=4, axis=1)
