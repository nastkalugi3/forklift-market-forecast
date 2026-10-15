"""СИНТЕТИЧЕСКИЙ тест пайплайна: копия файла заказчика + один выдуманный квартал справа.

  python tests/make_test_quarter.py  →  tests/_test_new_quarter.xlsx
  python update.py --xlsx tests/_test_new_quarter.xlsx     (лучше в копии проекта!)

Значения нового квартала = тот же квартал год назад × множитель (ДВС ×1,25, электро ×1,05), округлено до 10.
Это НЕ данные рынка — файл нужен только для проверки, что обновление проходит от начала до конца.
"""
import re
from pathlib import Path

import openpyxl
from openpyxl.utils import get_column_letter

ROOT = Path(__file__).resolve().parents[1]
src = ROOT / "data/raw/silant_forklift_market.xlsx"
out = ROOT / "tests/_test_new_quarter.xlsx"
wb = openpyxl.load_workbook(src)
ws = wb.worksheets[0]
QRE = re.compile(r"(\d)\s*кв\.?\s*(\d{4})")
for hr in [r for r in range(1, ws.max_row + 1) if any(QRE.search(str(c.value or "")) for c in ws[r])]:
    cols = [c.column for c in ws[hr] if QRE.search(str(c.value or ""))]
    last = max(cols); m = QRE.search(ws.cell(hr, last).value)
    qn, yr = int(m.group(1)), int(m.group(2))
    qn, yr = (1, yr + 1) if qn == 4 else (qn + 1, yr)
    new = last + 1
    ws.cell(hr, new, f"{qn} кв.{yr}")
    # колонка того же квартала год назад
    ly = [c for c in cols if QRE.search(ws.cell(hr, c).value).groups() == (str(qn), str(yr - 1))][0]
    r = hr + 1
    mult = 1.25 if "Мощность ДВС" in str(ws.cell(hr, 2).value) else 1.05
    first = r
    while ws.cell(r, 1).value and not str(ws.cell(r, 1).value).lower().startswith("итого"):
        ws.cell(r, new, int(round((ws.cell(r, ly).value or 0) * mult / 10) * 10)); r += 1
    L = get_column_letter(new)
    ws.cell(r, new, f"=SUM({L}{first}:{L}{r - 1})")
wb.save(out)
print("Синтетический тестовый файл:", out)
